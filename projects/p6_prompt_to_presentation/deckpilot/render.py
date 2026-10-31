"""PowerPoint renderer (python-pptx): payload JSON -> genuinely editable .pptx.

Everything on a slide is a native, editable object: text boxes (no placeholders, so
every viewer shows the text), native charts backed by embedded chart data (right-click >
Edit Data in PowerPoint), native tables, and speaker notes. No chart is an image.

Native threaded PowerPoint comments are not supported by python-pptx, so review comments
are exported into the speaker notes and a separate review log (CSV) instead.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = "0B2545"
INK = "1B2430"
GREY = "5C6773"
LIGHT = "F3F6FA"
LINE = "D5DCE4"
# Validated categorical order (CVD-checked); single-series charts use the template accent.
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
TONES = {"red": "B02A37", "amber": "C77D00", "green": "2D6A4F", "blue": "13315C", "grey": "8D99AE"}

W, H = 13.333, 7.5
LAYOUTS: dict[str, dict[str, Any]] = {
    "cards": {"slots": [(0.5, 1.55, 12.33, 5.05)], "comm": None},
    "steps": {"slots": [(0.5, 1.55, 12.33, 3.55)], "comm": (0.5, 5.3, 12.33, 1.35)},
    "kpis+2charts": {"slots": [(0.5, 1.5, 12.33, 1.15), (0.5, 2.8, 4.65, 3.85), (5.25, 2.8, 4.65, 3.85)],
                     "comm": (10.05, 2.8, 2.78, 3.85)},
    "2charts": {"slots": [(0.5, 1.5, 6.1, 3.75), (6.73, 1.5, 6.1, 3.75)], "comm": (0.5, 5.35, 12.33, 1.35)},
    "kpis+chart": {"slots": [(0.5, 1.5, 12.33, 1.15), (0.5, 2.8, 8.1, 3.85)], "comm": (8.8, 2.8, 4.03, 3.85)},
    "table+comm": {"slots": [(0.5, 1.5, 8.45, 5.15)], "comm": (9.15, 1.5, 3.68, 5.15)},
    "chart+comm": {"slots": [(0.5, 1.5, 8.45, 5.15)], "comm": (9.15, 1.5, 3.68, 5.15)},
    "chart+table": {"slots": [(0.5, 1.5, 6.0, 5.15), (6.7, 1.5, 6.13, 3.2)], "comm": (6.7, 4.85, 6.13, 1.8)},
    "table+chart": {"slots": [(0.5, 1.5, 12.33, 2.85), (0.5, 4.5, 6.0, 2.15)], "comm": (6.7, 4.5, 6.13, 2.15)},
    "chart+cards": {"slots": [(0.5, 1.5, 6.6, 5.15), (7.3, 1.5, 5.53, 3.35)], "comm": (7.3, 5.0, 5.53, 1.65)},
    "commentary": {"slots": [], "comm": (0.5, 1.55, 12.33, 5.1)},
}


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _box(slide, x, y, w, h, text="", size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = str(line)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)
        r.font.name = "Calibri"
    return tb


def _rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


# ---------------------------------------------------------------------- blocks
def draw_kpis(slide, block, x, y, w, h, accent):
    items = block["items"]
    n = max(len(items), 1)
    gap = 0.15
    tw = (w - gap * (n - 1)) / n
    for i, it in enumerate(items):
        tx = x + i * (tw + gap)
        _rect(slide, tx, y, tw, h, LIGHT, LINE, name=f"KPI {i + 1}")
        _rect(slide, tx, y, 0.07, h, TONES.get(it.get("tone", "grey"), accent))
        _box(slide, tx + 0.15, y + 0.05, tw - 0.2, h * 0.48, it["value"], size=22 if len(it["value"]) < 12 else 16,
             bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        _box(slide, tx + 0.15, y + h * 0.5, tw - 0.2, h * 0.24, it["label"], size=10, bold=True, color=INK)
        if it.get("sub"):
            _box(slide, tx + 0.15, y + h * 0.72, tw - 0.2, h * 0.27, it["sub"], size=8, color=GREY)


CHART_TYPES = {
    "line": XL_CHART_TYPE.LINE_MARKERS, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED, "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT, "waterfall": XL_CHART_TYPE.COLUMN_STACKED,
}


def _waterfall_series(categories, values):
    """Native stacked-column waterfall: invisible base + increase + decrease + totals (all editable)."""
    base, up, down, total = [], [], [], []
    running = 0.0
    last = len(values) - 1
    for i, v in enumerate(values):
        v = v or 0.0
        if i == 0 or i == last:
            base.append(0.0)
            up.append(0.0)
            down.append(0.0)
            total.append(v)
            running = v
        else:
            nxt = running + v
            base.append(min(running, nxt))
            up.append(v if v > 0 else 0.0)
            down.append(-v if v < 0 else 0.0)
            total.append(0.0)
            running = nxt
    return [("Base", base), ("Total", total), ("Increase", up), ("Decrease", down)]


def draw_chart(slide, block, x, y, w, h, accent):
    note = block.get("note")
    if note:
        h -= 0.35
    kind = block["chart"]
    cd = CategoryChartData()
    cd.categories = block["categories"]
    series = block["series"]
    if kind == "waterfall":
        series = [{"name": n, "values": v} for n, v in _waterfall_series(block["categories"], series[0]["values"])]
    for s in series:
        cd.add_series(s["name"], s["values"], number_format=block.get("number_format") or "General")
    gf = slide.shapes.add_chart(CHART_TYPES[kind], Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = f"Chart: {block['title'][:40]}"
    ch = gf.chart
    ch.font.size = Pt(9)
    ch.font.name = "Calibri"
    ch.has_title = True
    ch.chart_title.text_frame.text = block["title"]
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size = Pt(11)
    tp.runs[0].font.bold = True
    tp.runs[0].font.color.rgb = rgb(NAVY)
    multi = len(series) > 1 and kind != "waterfall"
    ch.has_legend = multi or kind == "doughnut"
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(9)
    plot = ch.plots[0]
    if kind in ("column", "bar", "bar_stacked", "column_stacked", "waterfall"):
        plot.gap_width = 60 if kind != "waterfall" else 40
        if kind in ("bar_stacked", "column_stacked", "waterfall"):
            plot.overlap = 100
    if kind == "waterfall":
        colors = [None, NAVY, "2D6A4F", "B02A37"]
        for s, c in zip(plot.series, colors, strict=True):
            if c is None:
                s.format.fill.background()
                s.format.line.fill.background()
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = rgb(c)
    elif kind == "doughnut":
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(PALETTE[i % len(PALETTE)])
        plot.has_data_labels = True
        plot.data_labels.number_format = "0%"
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(9)
    else:
        for i, s in enumerate(plot.series):
            col = PALETTE[i % len(PALETTE)] if multi else accent.lstrip("#")
            if kind == "line":
                s.format.line.color.rgb = rgb(col)
                s.format.line.width = Pt(2.25)
                s.smooth = False
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = rgb(col)
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = rgb(col)
        hl = set(block.get("highlight") or [])
        if hl and kind == "column" and len(plot.series) == 1:
            for i, c in enumerate(block["categories"]):
                pt = plot.series[0].points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(accent.lstrip("#") if c in hl else "AEB8C4")
        if block.get("data_labels"):
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(8)
            dl.number_format = block.get("number_format") or "General"
            dl.number_format_is_linked = False
            if kind in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
    if kind != "doughnut":
        try:
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb("E3E8EE")
            va.tick_labels.font.size = Pt(8)
            va.tick_labels.number_format = block.get("number_format") or "General"
            va.tick_labels.number_format_is_linked = False
            va.format.line.fill.background()
            ca = ch.category_axis
            ca.tick_labels.font.size = Pt(8)
            ca.format.line.color.rgb = rgb(LINE)
        except (ValueError, AttributeError):
            pass
    if note:
        _box(slide, x, y + h + 0.02, w, 0.33, note, size=8, color=GREY, italic=True)


def _heat(v: float, vmax: float) -> tuple[str, str]:
    t = max(0.0, min(1.0, v / vmax)) if vmax else 0.0
    lo, hi = (243, 246, 250), (19, 49, 92)
    c = "".join(f"{int(lo[i] + (hi[i] - lo[i]) * t):02X}" for i in range(3))
    return c, ("FFFFFF" if t > 0.5 else INK)


def draw_table(slide, block, x, y, w, h, accent):
    cols, rows = block["columns"], block["rows"]
    title_h = 0.3 if block.get("title") else 0
    if title_h:
        _box(slide, x, y, w, title_h, block["title"], size=11, bold=True, color=NAVY)
    nrows = len(rows) + 1
    row_h = min(0.36, max(0.22, (h - title_h) / max(nrows, 1)))
    shape = slide.shapes.add_table(nrows, len(cols), Inches(x), Inches(y + title_h), Inches(w),
                                   Inches(row_h * nrows))
    shape.name = f"Table: {block.get('title', '')[:40]}"
    tbl = shape.table
    widths = block.get("col_widths")
    if widths:
        scale = w / sum(widths)
        for i, cw in enumerate(widths):
            tbl.columns[i].width = Emu(int(Inches(cw * scale)))
    else:
        first = min(2.6, w * 0.32) if len(cols) > 3 else w / len(cols)
        rest = (w - first) / max(len(cols) - 1, 1)
        for i in range(len(cols)):
            tbl.columns[i].width = Emu(int(Inches(first if i == 0 else rest)))
    fs = 9 if len(cols) <= 6 else 8
    if nrows > 12:
        fs = 7
    heat = block.get("heat")
    for r in range(nrows):
        tbl.rows[r].height = Emu(int(Inches(row_h)))
        for c in range(len(cols)):
            cell = tbl.cell(r, c)
            val = cols[c] if r == 0 else rows[r - 1][c]
            txt = "" if val is None else (f"{val:,.1f}" if isinstance(val, float) else str(val))
            cell.text = txt
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            numeric = txt[:1].isdigit() or txt.startswith(("-", "+", "€", "₹"))
            p.alignment = PP_ALIGN.RIGHT if c > 0 and numeric else PP_ALIGN.LEFT
            run = p.runs[0] if p.runs else p.add_run()
            run.font.size = Pt(fs)
            run.font.name = "Calibri"
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(NAVY)
                run.font.bold = True
                run.font.color.rgb = rgb("FFFFFF")
            elif heat and c >= heat.get("from_col", 1) and isinstance(val, (int, float)):
                bg, fg = _heat(float(val), float(heat.get("max") or 1))
                cell.fill.fore_color.rgb = rgb(bg)
                run.font.color.rgb = rgb(fg)
            else:
                cell.fill.fore_color.rgb = rgb("FFFFFF" if r % 2 else LIGHT)
                run.font.color.rgb = rgb(INK)
                if txt in ("Worse", "RED", "QUARANTINE", "ERROR"):
                    run.font.color.rgb = rgb(TONES["red"])
                elif txt in ("Better", "OK"):
                    run.font.color.rgb = rgb(TONES["green"])
                elif txt in ("WARNING", "Flat"):
                    run.font.color.rgb = rgb(TONES["amber"])


def draw_cards(slide, block, x, y, w, h, accent, vertical=False):
    items = block["items"]
    n = len(items)
    if not n:
        return
    numbered = block.get("numbered")
    if vertical:
        gap = 0.1
        ch_ = (h - gap * (n - 1)) / n
        boxes = [(x, y + i * (ch_ + gap), w, ch_) for i in range(n)]
    elif n == 4 and not numbered:
        gap = 0.2
        cw, chh = (w - gap) / 2, (h - gap) / 2
        boxes = [(x + (i % 2) * (cw + gap), y + (i // 2) * (chh + gap), cw, chh) for i in range(4)]
    else:
        gap = 0.2
        cw = (w - gap * (n - 1)) / n
        h = min(h, 3.6)
        boxes = [(x + i * (cw + gap), y, cw, h) for i in range(n)]
    for i, (it, (bx, by, bw, bh)) in enumerate(zip(items, boxes, strict=True)):
        _rect(slide, bx, by, bw, bh, LIGHT, LINE, name=f"Card {i + 1}")
        _rect(slide, bx, by, bw, 0.06, accent.lstrip("#"))
        top = by + 0.15
        if numbered:
            c = _rect(slide, bx + 0.2, top, 0.55, 0.55, accent.lstrip("#"), shape=MSO_SHAPE.OVAL)
            c.text_frame.text = str(i + 1)
            pr = c.text_frame.paragraphs[0]
            pr.alignment = PP_ALIGN.CENTER
            pr.runs[0].font.size = Pt(18)
            pr.runs[0].font.bold = True
            pr.runs[0].font.color.rgb = rgb("FFFFFF")
            top += 0.7
        small = vertical and bh < 1.3
        _box(slide, bx + 0.2, top, bw - 0.4, 0.5 if not small else 0.32, it["head"], size=12 if small else 15,
             bold=True, color=NAVY)
        _box(slide, bx + 0.2, top + (0.55 if not small else 0.32), bw - 0.4, bh - (top - by) - 0.75, it["body"],
             size=9 if small else 12, color=INK)
        if it.get("when"):
            _box(slide, bx + 0.2, by + bh - 0.5, bw - 0.4, 0.35, it["when"], size=11, bold=True,
                 color=accent.lstrip("#"))


def draw_commentary(slide, items, x, y, w, h, accent, big=False, approved=False):
    _rect(slide, x, y, w, h, "FFFFFF", LINE, name="Management commentary")
    _rect(slide, x, y, 0.06, h, accent.lstrip("#"))
    head = "Management commentary" + ("" if approved else " (draft)")
    _box(slide, x + 0.15, y + 0.06, w - 0.25, 0.32, head, size=11, bold=True, color=NAVY)
    tb = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.4), Inches(w - 0.25), Inches(h - 0.45))
    tb.name = "Commentary text"
    tf = tb.text_frame
    tf.word_wrap = True
    size = 14 if big else (10 if h > 2 else 9)
    for i, c in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        kind = c.get("kind", "fact")
        prefix = {"hypothesis": "Hypothesis: ", "caveat": "Caveat: "}.get(kind, "")
        r = p.add_run()
        r.text = "• " + prefix + c["text"]
        r.font.size = Pt(size)
        r.font.name = "Calibri"
        r.font.italic = kind != "fact"
        r.font.color.rgb = rgb(GREY if kind == "caveat" else INK)


# ---------------------------------------------------------------------- slide chrome
def _chrome(slide, deck, s, no, total):
    meta = deck["meta"]
    accent = deck.get("accent", "#13315c").lstrip("#")
    _rect(slide, 0, 0, W, 0.95, NAVY, name="Header bar")
    _rect(slide, 0, 0.95, W, 0.05, accent, name="Accent stripe")
    _box(slide, 0.5, 0.12, W - 1.0, 0.75, s["title"], size=24 if len(s["title"]) < 60 else 18, bold=True,
         color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, name="Title")
    _box(slide, 0.5, 1.05, W - 1.0, 0.38, s.get("subtitle", ""), size=12, color=GREY, name="Subtitle")
    if s.get("footnote"):
        _box(slide, 0.5, 6.72, W - 1.0, 0.3, s["footnote"], size=8, color=GREY, italic=True, name="Footnote")
    _rect(slide, 0.5, 7.05, W - 1.0, 0.01, LINE)
    left = (f"{meta['org']} · {'Synthetic data' if meta['synthetic'] else 'Uploaded data'} · dataset "
            f"{meta['dataset_version']} · {meta['version_label']}")
    _box(slide, 0.5, 7.08, 8.5, 0.3, left, size=8, color=GREY, name="Footer")
    status = meta.get("status_text", "DRAFT – for internal review only")
    _box(slide, 9.0, 7.08, W - 9.5, 0.3, f"{status}  |  {no}", size=8, bold=True,
         color=TONES["green"] if meta.get("approved") else TONES["amber"], align=PP_ALIGN.RIGHT, name="Status")


def _cover(slide, deck, s):
    meta = deck["meta"]
    accent = deck.get("accent", "#13315c").lstrip("#")
    _rect(slide, 0, 0, W, H, NAVY, name="Background")
    _rect(slide, 0.6, 1.0, 0.1, 2.3, accent if accent != NAVY else "2A9D8F")
    _box(slide, 0.9, 0.85, 11.5, 0.4, deck.get("kicker", "").upper(), size=12, bold=True, color="9FB3C8")
    _box(slide, 0.9, 1.25, 11.5, 1.4, s["title"], size=34 if len(s["title"]) < 50 else 28, bold=True,
         color="FFFFFF", name="Title")
    _box(slide, 0.9, 2.7, 11.5, 0.5, s.get("subtitle", ""), size=14, color="D6E2F0", name="Subtitle")
    _box(slide, 0.9, 3.25, 11.5, 0.4, f"Generated {meta['generated_at']} by Prompt-to-Presentation AI · "
         f"{meta['version_label']}", size=10, color="9FB3C8")
    for b in s["blocks"]:
        if b["kind"] == "kpis":
            items = b["items"]
            n = len(items)
            tw = (11.5 - 0.3 * (n - 1)) / n
            for i, it in enumerate(items):
                tx = 0.9 + i * (tw + 0.3)
                _rect(slide, tx, 4.3, tw, 1.75, "13315C", name=f"KPI {i + 1}")
                _box(slide, tx + 0.2, 4.4, tw - 0.4, 0.8, it["value"], size=34, bold=True, color="FFFFFF")
                _box(slide, tx + 0.2, 5.2, tw - 0.4, 0.4, it["label"], size=12, bold=True, color="D6E2F0")
                _box(slide, tx + 0.2, 5.55, tw - 0.4, 0.4, it.get("sub", ""), size=10, color="9FB3C8")
    status = meta.get("status_text", "DRAFT – for internal review only")
    _box(slide, 0.9, 6.6, 11.5, 0.6, [f"{meta['org']} · {'Synthetic data' if meta['synthetic'] else 'Uploaded data'}"
                                      f" · dataset {meta['dataset_version']}", status], size=9, color="9FB3C8")


def notes_text(deck, s, no) -> str:
    meta = deck["meta"]
    parts = []
    talk = " ".join(c["text"] for c in s.get("commentary", []) if c.get("kind") == "fact")
    if talk:
        parts.append("TALK TRACK: " + talk)
    hyp = [c["text"] for c in s.get("commentary", []) if c.get("kind") == "hypothesis"]
    if hyp:
        parts.append("HYPOTHESES (not proven by the data): " + " ".join(hyp))
    ev = []
    for c in s.get("commentary", []):
        e = c.get("evidence") or {}
        if e.get("value") is not None:
            ev.append(f"- {e.get('metric')} = {e.get('value')} | window: {e.get('window')} | "
                      f"filters: {e.get('filters')}")
    if ev:
        parts.append("EVIDENCE:\n" + "\n".join(ev))
    if s.get("human_edited"):
        parts.append("NOTE: commentary on this slide was edited by a person"
                     + (f" ({s.get('edited_by')})" if s.get("edited_by") else "") + "; numbers in edited text are "
                     "not auto-verified" + (f" (unverified: {', '.join(s['unverified'])})" if s.get("unverified")
                                            else "") + ".")
    if s.get("notes_extra"):
        parts.append(s["notes_extra"])
    disc = meta.get("disclosures") or []
    if disc and no in (1, len(deck["slides"])):
        parts.append("DISCLOSURES: " + " ".join(disc[:6]))
    parts.append(f"PROVENANCE: dataset {meta['dataset_version']} (sha {meta['dataset_sha'][:12]}), period "
                 f"{meta['period']}, filters {meta['filters_label']}, "
                 f"queries {', '.join(meta.get('fingerprints', [])[:6])}"
                 f", parser {meta.get('parser_mode')}, generated {meta['generated_at']}.")
    comments = (deck.get("comments") or {}).get(str(no)) or []
    if comments:
        parts.append("REVIEWER COMMENTS:\n" + "\n".join(f"- [{c['status']}] {c['author']}: {c['text']}"
                                                        for c in comments))
    return "\n\n".join(parts)


def render(deck: dict[str, Any], path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    accent = deck.get("accent", "#13315c")
    prs.core_properties.title = deck["slides"][0]["title"][:250]
    prs.core_properties.author = "Prompt-to-Presentation AI (synthetic demo)"
    prs.core_properties.subject = f"{deck['meta']['version_label']} · {deck['meta']['dataset_version']}"
    prs.core_properties.keywords = f"deck_id={deck['meta'].get('deck_id')}; status={deck['meta'].get('status')}"
    total = len(deck["slides"])
    for no, s in enumerate(deck["slides"], start=1):
        slide = prs.slides.add_slide(blank)
        if s["layout"] == "cover":
            _cover(slide, deck, s)
        else:
            _chrome(slide, deck, s, no, total)
            lay = LAYOUTS[s["layout"]]
            for block, (x, y, w, h) in zip(s["blocks"], lay["slots"], strict=False):
                k = block["kind"]
                if k == "kpis":
                    draw_kpis(slide, block, x, y, w, h, accent)
                elif k == "chart":
                    draw_chart(slide, block, x, y, w, h, accent)
                elif k == "table":
                    draw_table(slide, block, x, y, w, h, accent)
                elif k == "cards":
                    draw_cards(slide, block, x, y, w, h, accent, vertical=s["layout"] == "chart+cards")
            if lay["comm"] and s.get("commentary"):
                draw_commentary(slide, s["commentary"], *lay["comm"], accent, big=s["layout"] == "commentary",
                                approved=bool(deck["meta"].get("approved")))
        if deck["meta"].get("speaker_notes", True):
            slide.notes_slide.notes_text_frame.text = notes_text(deck, s, no)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(f"{path} exists; versions are never overwritten")
    prs.save(str(path))
    return path


def to_pdf(pptx_path: str | Path, out_dir: str | Path, timeout: int = 120) -> Path | None:
    """Convert with LibreOffice if it is installed; returns None when PDF export is unavailable."""
    import shutil
    import subprocess

    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)],
                       check=True, capture_output=True, timeout=timeout)
    except (subprocess.SubprocessError, OSError):
        return None
    pdf = out_dir / (Path(pptx_path).stem + ".pdf")
    return pdf if pdf.exists() else None
