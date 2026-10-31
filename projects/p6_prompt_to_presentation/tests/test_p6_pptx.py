"""Generated decks are valid, editable PowerPoint files."""

from pathlib import Path

import pytest
from p6_helpers import APAC_PROMPT
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckpilot.render import to_pdf


def shapes(path):
    prs = Presentation(path)
    return prs, [list(s.shapes) for s in prs.slides]


def test_six_slide_deck_is_native_and_editable(service):
    rec, _ = service.generate(APAC_PROMPT, "ops_risk", actor="maker")
    prs, sl = shapes(rec.file_path)
    assert len(prs.slides) == 6
    all_shapes = [sh for s in sl for sh in s]
    assert not [sh for sh in all_shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]  # no image charts
    assert sum(1 for sh in all_shapes if sh.has_chart) >= 3
    assert sum(1 for sh in all_shapes if sh.has_table) >= 2
    for s in prs.slides:
        notes = s.notes_slide.notes_text_frame.text
        assert "PROVENANCE: dataset ops-" in notes
    texts = " ".join(sh.text_frame.text for sh in all_shapes if sh.has_text_frame)
    assert "DRAFT" in texts and "Synthetic data" in texts
    chart = next(sh for sh in all_shapes if sh.has_chart).chart
    assert len(chart.plots[0].categories) > 0  # embedded, editable chart data


def test_full_workforce_pack_renders_every_slide_type(service):
    rec, _ = service.generate("Create the full workforce review pack for 2018 with speaker notes", "workforce",
                              actor="maker")
    prs, _ = shapes(rec.file_path)
    assert len(prs.slides) == 11
    for extra in ["add a profession mix slide", "add data quality", "add management commentary"]:
        rec, _ = service.follow_up(rec.version_id, extra, actor="maker")
    prs, _ = shapes(rec.file_path)
    assert len(prs.slides) == 12  # capped at the 12-slide maximum


def test_every_ops_slide_type_renders(service):
    d = service.domains["ops_risk"]
    rec, _ = service.generate("EMEA risk pack Q3 2026 with 12 slides, ageing, remediation, loss, high-severity "
                              "exceptions, data quality, operations performance, trend", "ops_risk", actor="maker")
    types = {s["slide_type"] for s in rec.payload["slides"]}
    assert len(types) == 12 and types <= set(d.slide_catalog)
    assert Presentation(rec.file_path).slides


def test_versions_never_overwrite(service, tmp_path):
    from deckpilot.render import render

    rec, _ = service.generate(APAC_PROMPT, "ops_risk", actor="maker")
    with pytest.raises(FileExistsError):
        render(rec.payload, rec.file_path)


def test_pdf_export_when_libreoffice_available(service):
    import shutil

    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice not installed; PDF export is reported as unavailable in the app")
    rec, _ = service.generate(APAC_PROMPT, "ops_risk", actor="maker")
    pdf = to_pdf(rec.file_path, Path(rec.file_path).parent)
    assert pdf and pdf.read_bytes()[:4] == b"%PDF"
